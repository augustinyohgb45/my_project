"""
Import du catalogue depuis le classeur Excel « gestion_Stock Dynasty.xlsm ».

Lit la feuille « Catalogue » (ID_Produit, Designation, Categorie, Prix_Achat, Prix_Vente,
Stock_Alerte) et le stock initial de la feuille « Stock_Reel ». Les données sont nettoyées au
passage : espaces superflus, catégories harmonisées, prix d'achat incohérents écartés.
Un nouvel import met à jour les produits existants (même code) sans dupliquer le stock initial.
"""
import re
from dataclasses import dataclass, field

from django.db import transaction
from openpyxl import load_workbook

from .models import Categorie, MouvementStock, Produit

# Ordre d'affichage des catégories (les autres viennent ensuite)
ORDRE_CATEGORIES = [
    "Champagne", "Vin mousseux", "Vin", "Whisky", "Cognac", "Vodka", "Tequila", "Gin", "Rhum",
    "Liqueur", "Apéritif", "Bière", "Soft", "Eau",
]

# Orthographes de l'Excel ramenées à une seule catégorie
SYNONYMES_CATEGORIES = {"bières": "Bière", "biere": "Bière", "bieres": "Bière"}

# La catégorie « Liqueur » de l'Excel regroupe tous les alcools forts : on les répartit d'après le nom
TYPES_ALCOOLS_FORTS = [
    ("hennessy", "Cognac"),
    ("johnnie", "Whisky"), ("jack daniel", "Whisky"), ("chivas", "Whisky"),
    ("ballantines", "Whisky"), ("ballantine", "Whisky"), ("jb", "Whisky"),
    ("gordon", "Gin"), ("belvedere", "Vodka"), ("saint james", "Rhum"),
    ("don julio", "Tequila"), ("casamigos", "Tequila"), ("azul", "Tequila"),
    ("baileys", "Liqueur"), ("jägermeister", "Liqueur"), ("martini", "Apéritif"),
]

# En dessous de 10 % du prix de vente, le prix d'achat est considéré comme une erreur de saisie
RATIO_PRIX_ACHAT_MINIMUM = 0.10

COLONNES_CATALOGUE = ('ID_Produit', 'Designation', 'Categorie', 'Prix_Achat', 'Prix_Vente', 'Stock_Alerte')


class ErreurImport(Exception):
    pass


@dataclass
class RapportImport:
    crees: list = field(default_factory=list)
    mis_a_jour: list = field(default_factory=list)
    stocks_initiaux: list = field(default_factory=list)
    avertissements: list = field(default_factory=list)
    sans_prix_achat: list = field(default_factory=list)
    sans_seuil: list = field(default_factory=list)

    @property
    def nb_produits(self):
        return len(self.crees) + len(self.mis_a_jour)


def nettoyer_texte(valeur):
    return re.sub(r"\s+", " ", str(valeur or "")).strip()


def _entier(valeur):
    if valeur is None or valeur == '':
        return None
    try:
        return int(round(float(valeur)))
    except (TypeError, ValueError):
        return None


def normaliser_categorie(categorie, nom):
    categorie = nettoyer_texte(categorie)
    categorie = SYNONYMES_CATEGORIES.get(categorie.lower(), categorie)
    if categorie.lower() == "liqueur":
        nom_minuscule = nom.lower()
        for mot, type_alcool in TYPES_ALCOOLS_FORTS:
            if re.search(rf"\b{re.escape(mot)}\b", nom_minuscule):
                return type_alcool
    return categorie[:1].upper() + categorie[1:] if categorie else "Autre"


def _obtenir_categorie(nom, cache):
    if nom not in cache:
        ordre = ORDRE_CATEGORIES.index(nom) + 1 if nom in ORDRE_CATEGORIES else 100
        cache[nom], _ = Categorie.objects.get_or_create(nom=nom, defaults={'ordre': ordre})
    return cache[nom]


def lire_classeur(fichier):
    """Retourne (lignes du catalogue, stock initial par code) lus dans le classeur."""
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception as e:
        raise ErreurImport(f"Fichier Excel illisible : {e}")
    try:
        if 'Catalogue' not in classeur.sheetnames:
            raise ErreurImport("La feuille « Catalogue » est introuvable dans ce classeur.")
        lignes = list(classeur['Catalogue'].iter_rows(values_only=True))
        if not lignes:
            raise ErreurImport("La feuille « Catalogue » est vide.")
        entetes = [nettoyer_texte(e) for e in lignes[0]]
        manquantes = [c for c in COLONNES_CATALOGUE if c not in entetes]
        if manquantes:
            raise ErreurImport(f"Colonnes manquantes dans « Catalogue » : {', '.join(manquantes)}.")
        index = {c: entetes.index(c) for c in COLONNES_CATALOGUE}
        catalogue = []
        for numero, ligne in enumerate(lignes[1:], start=2):
            valeurs = {c: (ligne[i] if i < len(ligne) else None) for c, i in index.items()}
            if not any(v not in (None, '') for v in valeurs.values()):
                continue
            catalogue.append((numero, valeurs))

        # Stock_Reel : la colonne A reprend le code du catalogue sur la même ligne (=Catalogue!A2…)
        stocks = {}
        if 'Stock_Reel' in classeur.sheetnames:
            codes_par_ligne = {numero: nettoyer_texte(v['ID_Produit']) for numero, v in catalogue}
            for numero, ligne in enumerate(classeur['Stock_Reel'].iter_rows(min_row=2, values_only=True), start=2):
                if len(ligne) < 3:
                    continue
                code = nettoyer_texte(ligne[0]) or codes_par_ligne.get(numero, '')
                quantite = _entier(ligne[2])
                if code and quantite:
                    stocks[code] = quantite
        return catalogue, stocks
    finally:
        classeur.close()


def importer_catalogue(fichier, utilisateur=None, simulation=False):
    """Importe le catalogue du classeur. En simulation, rien n'est enregistré."""
    catalogue, stocks = lire_classeur(fichier)
    rapport = RapportImport()
    categories = {}
    codes_vus = set()

    with transaction.atomic():
        for numero, valeurs in catalogue:
            code = nettoyer_texte(valeurs['ID_Produit'])
            nom = nettoyer_texte(valeurs['Designation'])
            if not code or not nom:
                rapport.avertissements.append(f"Ligne {numero} ignorée : code ou désignation manquant.")
                continue
            if code in codes_vus:
                rapport.avertissements.append(f"Ligne {numero} ignorée : le code {code} apparaît deux fois.")
                continue
            codes_vus.add(code)

            prix_vente = _entier(valeurs['Prix_Vente'])
            if prix_vente is None:
                rapport.avertissements.append(f"Ligne {numero} ignorée ({nom}) : prix de vente manquant.")
                continue
            prix_achat = _entier(valeurs['Prix_Achat'])
            if prix_achat is not None and prix_achat < prix_vente * RATIO_PRIX_ACHAT_MINIMUM:
                rapport.avertissements.append(
                    f"{nom} : prix d'achat de {prix_achat} FCFA écarté (incohérent avec un prix de vente "
                    f"de {prix_vente} FCFA). À ressaisir.")
                prix_achat = None
            seuil = _entier(valeurs['Stock_Alerte']) or 0

            categorie = _obtenir_categorie(normaliser_categorie(valeurs['Categorie'], nom), categories)
            produit, cree = Produit.objects.update_or_create(code=code, defaults={
                'nom': nom, 'categorie': categorie, 'prix_achat': prix_achat,
                'prix_vente': prix_vente, 'seuil_alerte': seuil,
            })
            (rapport.crees if cree else rapport.mis_a_jour).append(produit)
            if prix_achat is None:
                rapport.sans_prix_achat.append(produit)
            if not seuil:
                rapport.sans_seuil.append(produit)

            # Stock initial : uniquement pour un produit qui n'a encore aucun mouvement
            quantite = stocks.get(code)
            if quantite and not produit.mouvements.exists():
                MouvementStock.objects.create(
                    produit=produit, type=MouvementStock.INVENTAIRE, quantite=quantite,
                    prix_unitaire=prix_achat, motif="Stock initial (import Excel)", utilisateur=utilisateur)
                rapport.stocks_initiaux.append((produit, quantite))

        if simulation:
            transaction.set_rollback(True)
    return rapport
